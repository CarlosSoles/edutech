import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from fpdf import FPDF

from docx.enum.section import WD_ORIENT

def generate_word_report(header_data, students_data):
    doc = Document()
    
    # Orientación Horizontal (Landscape)
    section = doc.sections[0]
    new_width, new_height = section.page_height, section.page_width
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = new_width
    section.page_height = new_height
    
    # Estilos globales
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Títulos
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p1.add_run("INSTRUMENTO DE EVALUACIÓN\nCUADERNO DE CAMPO (REGISTRO DE EVIDENCIAS)")
    run1.bold = True
    run1.underline = True
    
    doc.add_paragraph() # Espacio
    
    # Tabla 1: Cabecera Completa (4 filas x 3 columnas)
    table1 = doc.add_table(rows=4, cols=3)
    table1.style = 'Table Grid'
    
    r0 = table1.rows[0].cells
    r0[0].text = "Título de la Actividad de Aprendizaje"
    r0[1].text = header_data.get("titulo", "")
    r0[2].text = f"Fecha: {header_data.get('fecha', '')}"
    
    r1 = table1.rows[1].cells
    r1[0].text = "AULA:"
    r1[1].text = header_data.get("aula", "")
    r1[2].text = f"EDAD: {header_data.get('edad', '')} AÑOS"
    
    r2 = table1.rows[2].cells
    r2[0].text = f"ÁREA:\n{header_data.get('area', '')}"
    r2[1].text = f"COMPETENCIA:\n{header_data.get('competencia', '')}"
    r2[2].text = f"ESTÁNDAR DE APRENDIZAJE:\n{header_data.get('estandar', '')}"
    
    r3 = table1.rows[3].cells
    r3[0].text = "CRITERIOS DE EVALUACIÓN"
    r3[1].text = header_data.get("criterios", "")
    r3[2].text = f"CAPACIDADES:\n{header_data.get('capacidades', '')}"
    
    doc.add_paragraph() # Espacio
    
    # Tabla 3: Evidencias
    table3 = doc.add_table(rows=1 + len(students_data), cols=3)
    table3.style = 'Table Grid'
    
    h_cells = table3.rows[0].cells
    h_cells[0].text = "EVIDENCIAS / RETROALIMENTACIÓN\n\nRELACIÓN DE NIÑOS Y NIÑAS"
    h_cells[1].text = "DESCRIPCIÓN DE LAS EVIDENCIAS"
    h_cells[2].text = "ASPECTOS A RETROALIMENTAR"
    
    for i, child in enumerate(students_data):
        cells = table3.rows[i+1].cells
        cells[0].text = f"{i+1}. {child.get('nombre', '')}"
        cells[1].text = child.get("descripcion", "")
        cells[2].text = child.get("retroalimentacion", "")
    
    # Guardar en memoria
    byte_io = io.BytesIO()
    doc.save(byte_io)
    byte_io.seek(0)
    return byte_io.getvalue()
